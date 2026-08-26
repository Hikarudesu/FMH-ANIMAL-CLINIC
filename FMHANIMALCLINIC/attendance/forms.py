"""Forms for attendance and biometrics management."""
from django import forms
from django.core.exceptions import ValidationError
import csv
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from openpyxl import load_workbook
import xlrd
from .models import (
    BiometricDevice,
    DailyAttendance,
)


class BiometricDeviceForm(forms.ModelForm):
    """Form for creating/editing biometric devices."""
    
    class Meta:
        model = BiometricDevice
        fields = [
            'branch',
            'device_name',
            'device_model',
            'device_serial',
            'connection_type',
            'status',
            'is_active',
            'notes',
        ]
        widgets = {
            'branch': forms.Select(attrs={'class': 'form-control'}),
            'device_name': forms.TextInput(attrs={
                'class': 'form-control',
                'placeholder': 'e.g., Main Entrance Scanner',
            }),
            'device_model': forms.TextInput(attrs={
                'class': 'form-control',
                'placeholder': 'e.g., ZKTeco K20',
            }),
            'device_serial': forms.TextInput(attrs={
                'class': 'form-control',
                'placeholder': 'Serial number or unique ID',
            }),
            'connection_type': forms.Select(attrs={'class': 'form-control'}),
            'status': forms.Select(attrs={'class': 'form-control'}),
            'is_active': forms.CheckboxInput(attrs={'class': 'form-check-input'}),
            'notes': forms.Textarea(attrs={
                'class': 'form-control',
                'rows': 3,
                'placeholder': 'Location notes, setup details, etc.',
            }),
        }

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields['connection_type'].choices = [
            ('USB', 'USB / SD card file export'),
            ('MANUAL', 'CSV / Excel import'),
        ]
        self.fields['connection_type'].initial = 'MANUAL'
        self.fields['connection_type'].help_text = (
            'The scanner remains independent. Attendance is imported from its exported file.'
        )


class AttendanceImportForm(forms.Form):
    """Import one monthly attendance summary for payroll."""

    import_file = forms.FileField(
        widget=forms.FileInput(attrs={
            'class': 'form-control',
            'accept': '.csv,.xls,.xlsx,.ods,.xml,.xlsm',
        }),
        label='Import File',
        help_text='Upload one monthly summary export in CSV, XLS, XLSX, ODS, XLSM, or XML format.',
    )

    def clean(self):
        return super().clean()

    def clean_import_file(self):
        """Validate the import file."""
        file = self.cleaned_data.get('import_file')
        if not file:
            raise ValidationError('Please select a file to import.')

        # Check file size (max 10MB)
        if file.size > 10 * 1024 * 1024:
            raise ValidationError('File size exceeds 10MB limit.')

        # Check file extension
        filename = file.name.lower()
        allowed = ('.csv', '.xls', '.xlsx', '.ods', '.xlsm', '.xml')
        if not any(filename.endswith(ext) for ext in allowed):
            raise ValidationError('Only CSV, XLS, XLSX, ODS, XLSM, and XML spreadsheet files are supported.')

        return file

    @staticmethod
    def _normalise_header(value):
        if value is None:
            return ''
        return re.sub(r'[^a-z0-9]+', ' ', str(value).strip().lower()).strip()

    @staticmethod
    def _extract_cell_value(cell):
        if cell is None:
            return None
        if hasattr(cell, 'value'):
            return cell.value
        return cell

    def parse_csv(self, file):
        """Parse CSV file and return list of records."""
        file.seek(0)
        decoded_file = file.read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(decoded_file))
        return list(reader)

    def parse_excel(self, file):
        """Parse Excel file and return list of records."""
        file.seek(0)
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        records = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            record = {}
            for i, header in enumerate(headers):
                if header:
                    record[str(header)] = row[i] if i < len(row) else None
            if any(value is not None for value in record.values()):
                records.append(record)

        return records

    def parse_ods(self, file):
        """Parse an OpenDocument Spreadsheet (.ods) into row dictionaries."""
        file.seek(0)
        with zipfile.ZipFile(file) as archive:
            content = archive.read('content.xml')
        root = ET.fromstring(content)
        rows = []
        row_tag = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}table-row'
        cell_tag = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}table-cell'
        text_tag = '{urn:oasis:names:tc:opendocument:xmlns:text:1.0}p'
        repeat_attribute = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}number-columns-repeated'

        for row in root.iter(row_tag):
            values = []
            for cell in row.findall(cell_tag):
                value = ' '.join(text.text or '' for text in cell.iter(text_tag)).strip()
                repeat = int(cell.get(repeat_attribute, '1'))
                values.extend([value] * repeat)
            if any(value for value in values):
                rows.append(values)

        if not rows:
            return []
        headers = [str(value).strip() for value in rows[0]]
        return [
            {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            for values in rows[1:]
            if any(value for value in values)
        ]

    def parse_xls(self, file):
        """Parse legacy Excel .xls exports from older scanner software or summary spreadsheets."""
        file.seek(0)
        workbook = xlrd.open_workbook(file_contents=file.read())
        sheet = workbook.sheet_by_index(0)
        headers = [str(value).strip() for value in sheet.row_values(0)]
        records = []

        for row_index in range(1, sheet.nrows):
            values = []
            for cell in sheet.row(row_index):
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, workbook.datemode)
                values.append(value)
            record = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            if any(value is not None for value in record.values()):
                records.append(record)

        return records

    def parse_xml(self, file):
        """Parse monthly attendance summary blocks from Excel XML Spreadsheet exports."""
        file.seek(0)
        xml_text = file.read().decode('utf-8-sig', errors='ignore')
        root = ET.fromstring(xml_text)

        rows = []
        for row in root.iter():
            if row.tag.rsplit('}', 1)[-1] != 'Row':
                continue
            values = []
            for cell in row:
                if cell.tag.rsplit('}', 1)[-1] != 'Cell':
                    continue
                data = None
                for child in cell:
                    if child.tag.rsplit('}', 1)[-1] == 'Data':
                        data = child
                        break
                value = data.text if data is not None and data.text is not None else ''
                values.append(value)
            if values:
                rows.append(values)
        records = []
        for index, values in enumerate(rows):
            joined = ' '.join(str(value).strip() for value in values if value)
            if 'Name:' not in joined or 'ID:' not in joined or 'Date:' not in joined:
                continue

            identifier = re.search(r'ID\s*:\s*([^\s]+)', joined)
            period = re.search(
                r'Date\s*:\s*(\d{2})\.(\d{2})\.(\d{2}).*?(\d{2})\.(\d{2})\.(\d{2})',
                joined,
            )
            if not identifier or not period:
                continue

            start_year, start_month, start_day, end_year, end_month, end_day = period.groups()
            summary = {
                'Summary': 'MONTHLY',
                'Biometric ID': identifier.group(1),
                'Report Start': f'20{start_year}-{start_month}-{start_day}',
                'Report End': f'20{end_year}-{end_month}-{end_day}',
            }
            for detail_row in rows[index + 1:index + 3]:
                for value in detail_row:
                    match = re.match(r'\s*([^:：]+)[:：]\s*(.*?)\s*$', str(value or ''))
                    if match:
                        label, amount = match.groups()
                        summary[label.strip()] = amount.strip()
            records.append(summary)

        unique_records = {}
        for record in records:
            key = (record['Biometric ID'], record['Report Start'])
            unique_records[key] = record
        records = list(unique_records.values())
        if records:
            return records

        if not rows:
            return []

        headers = [str(value).strip() for value in rows[0]]
        for values in rows[1:]:
            record = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            if any(value is not None for value in record.values()):
                records.append(record)
        return records

    def get_records(self):
        """Parse the uploaded file based on its extension."""
        import_file = self.cleaned_data.get('import_file')
        filename = import_file.name.lower()

        if filename.endswith('.csv'):
            return self.parse_csv(import_file)
        if filename.endswith(('.xlsx', '.xlsm')):
            return self.parse_excel(import_file)
        if filename.endswith('.ods'):
            return self.parse_ods(import_file)
        if filename.endswith('.xls'):
            import_file.seek(0)
            signature = import_file.read(256).lstrip()
            import_file.seek(0)
            if signature.startswith(b'<?xml') or b'<Workbook' in signature:
                return self.parse_xml(import_file)
            return self.parse_xls(import_file)
        if filename.endswith('.xml'):
            return self.parse_xml(import_file)
        raise ValidationError('Unsupported attendance file format.')

class DailyAttendanceForm(forms.ModelForm):
    """Form for manually editing daily attendance records."""
    
    class Meta:
        model = DailyAttendance
        fields = [
            'check_in',
            'check_out',
            'status',
            'is_present',
            'adjustment_notes',
        ]
        widgets = {
            'check_in': forms.TimeInput(attrs={
                'class': 'form-control',
                'type': 'time',
            }),
            'check_out': forms.TimeInput(attrs={
                'class': 'form-control',
                'type': 'time',
            }),
            'status': forms.Select(attrs={'class': 'form-control'}),
            'is_present': forms.CheckboxInput(attrs={'class': 'form-check-input'}),
            'adjustment_notes': forms.Textarea(attrs={
                'class': 'form-control',
                'rows': 3,
                'placeholder': 'Notes about why this was adjusted',
            }),
        }
class AttendanceFilterForm(forms.Form):
    """Form for filtering attendance records."""
    
    staff = forms.ModelChoiceField(
        queryset=None,
        required=False,
        widget=forms.Select(attrs={'class': 'form-control'}),
        label='Staff Member',
    )
    
    device = forms.ModelChoiceField(
        queryset=BiometricDevice.objects.filter(is_active=True),
        required=False,
        widget=forms.Select(attrs={'class': 'form-control'}),
        label='Device',
    )
    
    date_from = forms.DateField(
        required=False,
        widget=forms.DateInput(attrs={
            'class': 'form-control',
            'type': 'date',
        }),
        label='From Date',
    )
    
    date_to = forms.DateField(
        required=False,
        widget=forms.DateInput(attrs={
            'class': 'form-control',
            'type': 'date',
        }),
        label='To Date',
    )
    
    status = forms.MultipleChoiceField(
        required=False,
        choices=DailyAttendance.AttendanceStatus.choices,
        widget=forms.CheckboxSelectMultiple(attrs={'class': 'form-check-input'}),
        label='Status',
    )
    
    def __init__(self, *args, **kwargs):
        """Initialize with dynamic queryset for staff."""
        super().__init__(*args, **kwargs)
        from employees.models import StaffMember
        self.fields['staff'].queryset = StaffMember.objects.filter(
            is_active=True
        ).order_by('first_name', 'last_name')
